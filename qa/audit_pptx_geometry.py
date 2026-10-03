#!/usr/bin/env python3
from __future__ import annotations
import argparse, json
from pathlib import Path
from pptx import Presentation

EMU_PER_INCH = 914400

def inch(v): return v / EMU_PER_INCH

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--json", dest="json_path", type=Path)
    args = ap.parse_args()
    prs = Presentation(args.pptx)
    sw, sh = inch(prs.slide_width), inch(prs.slide_height)
    failures=[]; slide_rows=[]
    for si, slide in enumerate(prs.slides, 1):
        outside=[]; tiny_text=[]
        for shape in slide.shapes:
            x,y,w,h = map(inch,(shape.left,shape.top,shape.width,shape.height))
            if x < -0.001 or y < -0.001 or x+w > sw+0.001 or y+h > sh+0.001:
                outside.append({"name":shape.name,"x":x,"y":y,"w":w,"h":h})
            if getattr(shape,"has_text_frame",False) and shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt < 8.5:
                            tiny_text.append({"name":shape.name,"pt":r.font.size.pt,"text":r.text[:80]})
        status='PASS' if not outside and not tiny_text else 'FAIL'
        slide_rows.append({"slide":si,"status":status,"outside":outside,"textBelow8_5pt":tiny_text})
        if status=='FAIL': failures.append(si)
    result={"pptx":str(args.pptx),"canvasIn":{"w":sw,"h":sh},"slides":slide_rows,"failedSlides":failures,"status":"PASS" if not failures else "FAIL"}
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if args.json_path: args.json_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if not failures else 1

if __name__=='__main__':
    raise SystemExit(main())
