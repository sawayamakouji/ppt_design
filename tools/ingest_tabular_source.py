#!/usr/bin/env python3
import csv, json, re, sys
from pathlib import Path
from datetime import datetime, date

def infer_scalar(v):
    if v is None or v == "": return None
    if isinstance(v, (int,float,bool,datetime,date)): return v
    s=str(v).strip()
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", s):
        try: return float(s) if "." in s else int(s)
        except: pass
    for f in ("%Y-%m-%d","%Y/%m/%d","%Y-%m","%Y/%m"):
        try: return datetime.strptime(s,f).date().isoformat()
        except: pass
    return s

def infer_col(name, vals):
    non=[v for v in vals if v not in (None,"")]
    typ="string"
    if non and all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in non): typ="number"
    elif non and all(isinstance(v,str) and re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?",v) for v in non): typ="date"
    n=name.lower()
    if typ=="date" or re.search(r"date|month|week|年月|月|日",name,re.I): sem="time"
    elif re.search(r"(^id$|_id$|code|コード|jan|sku)",n,re.I): sem="id"
    elif typ=="number": sem="metric"
    else: sem="dimension"
    return {"name":name,"label":name,"type":typ,"semantic":sem,"aggregation":"avg" if typ=="number" else "none"}

def load_rows(path: Path, sheet=None):
    suf=path.suffix.lower()
    if suf==".csv":
        text=path.read_text(encoding="utf-8-sig")
        r=csv.DictReader(text.splitlines())
        return [{k:infer_scalar(v) for k,v in row.items()} for row in r]
    if suf in (".xlsx",".xlsm"):
        from openpyxl import load_workbook
        wb=load_workbook(path, data_only=True, read_only=True)
        ws=wb[sheet] if sheet else wb[wb.sheetnames[0]]
        it=ws.iter_rows(values_only=True)
        headers=[str(x).strip() if x is not None else "" for x in next(it)]
        rows=[]
        for rr in it:
            d={headers[i]:infer_scalar(rr[i]) for i in range(min(len(headers),len(rr))) if headers[i]}
            if any(v not in (None,"") for v in d.values()): rows.append(d)
        return rows
    if suf==".json":
        data=json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data,dict) and "rows" in data: data=data["rows"]
        if not isinstance(data,list): raise ValueError("JSON must be an array of row objects or {rows:[...]}")
        return [{k:infer_scalar(v) for k,v in row.items()} for row in data]
    raise ValueError(f"unsupported input: {suf}")

def main():
    if len(sys.argv)<2:
        print("usage: ingest_tabular_source.py input.(csv|xlsx|json) [output.json] [table_id] [sheet]", file=sys.stderr); raise SystemExit(2)
    path=Path(sys.argv[1]); out=Path(sys.argv[2]) if len(sys.argv)>2 else Path("source-bundle.v1.json")
    tid=sys.argv[3] if len(sys.argv)>3 else path.stem
    sheet=sys.argv[4] if len(sys.argv)>4 else None
    rows=load_rows(path,sheet)
    names=list(rows[0].keys()) if rows else []
    cols=[infer_col(n,[r.get(n) for r in rows]) for n in names]
    bundle={"version":"1.0","profile":"SOURCE-BUNDLE-v1","metadata":{"title":path.name},"facts":[],"tables":[{"id":tid,"title":path.stem,"grain":"row","columns":cols,"rows":rows,"provenance":{"kind":path.suffix.lower().lstrip('.'),"ref":str(path)},"sample":False}]}
    out.write_text(json.dumps(bundle,ensure_ascii=False,indent=2,default=str),encoding="utf-8");print(out)
if __name__=="__main__": main()
